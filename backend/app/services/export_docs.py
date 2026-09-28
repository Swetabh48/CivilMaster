"""CivilMaster PDF / DOCX / DXF exports — assignment-readable layout."""

from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path
from typing import Any

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Flowable,
    HRFlowable,
    Image as RLImage,
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


INK = colors.HexColor("#0f172a")
MUTED = colors.HexColor("#475569")
ACCENT = colors.HexColor("#1d4e89")
Q_BG = colors.HexColor("#eef4fb")
A_BG = colors.HexColor("#e8f5ef")
A_BORDER = colors.HexColor("#2f6b4f")
LINE = colors.HexColor("#cbd5e1")
STEP_BG = colors.HexColor("#f8fafc")


class _Banner(Flowable):
    """Full-width colored label strip."""

    def __init__(self, text: str, bg: colors.Color, fg: colors.Color = colors.white, height: float = 22):
        super().__init__()
        self.text = text
        self.bg = bg
        self.fg = fg
        self._h = height
        self.width = 0

    def wrap(self, availWidth, availHeight):
        self.width = availWidth
        return availWidth, self._h

    def draw(self):
        self.canv.setFillColor(self.bg)
        self.canv.roundRect(0, 0, self.width, self._h, 3, fill=1, stroke=0)
        self.canv.setFillColor(self.fg)
        self.canv.setFont("Helvetica-Bold", 10)
        self.canv.drawString(8, 7, self.text)


def _clean(text: str) -> str:
    text = text or ""
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"`([^`]+)`", r"<font face='Courier' size='9'>\1</font>", text)
    text = re.sub(r"^#+\s*", "", text)
    return text


def _strip_md(line: str) -> str:
    line = (line or "").strip()
    line = re.sub(r"^#+\s*", "", line)
    line = re.sub(r"^\*\*(.+)\*\*$", r"\1", line)
    return line


def _png_from_pillow_fallback(
    *,
    title: str,
    lines: list[str],
    width: int = 900,
    height: int = 480,
) -> bytes:
    """Always-available diagram when SVG rasterizers are missing."""
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (width, height), "#f7f6f2")
    draw = ImageDraw.Draw(img)
    try:
        font_title = ImageFont.truetype("arial.ttf", 28)
        font = ImageFont.truetype("arial.ttf", 20)
    except Exception:
        font_title = ImageFont.load_default()
        font = font_title

    draw.rectangle([20, 20, width - 20, height - 20], outline="#1c2430", width=3)
    draw.text((40, 36), title[:60], fill="#1c2430", font=font_title)
    y = 90
    for line in lines[:8]:
        draw.text((40, y), line[:70], fill="#2f5d8a", font=font)
        y += 36

    # Simple beam glyph
    draw.line([(80, 340), (820, 340)], fill="#1c2430", width=5)
    draw.polygon([(80, 340), (60, 370), (100, 370)], fill="#1c2430")
    draw.rectangle([800, 340, 840, 365], outline="#1c2430", width=3)
    for x in range(120, 800, 80):
        draw.line([(x, 290), (x, 335)], fill="#c4a35a", width=2)
    draw.line([(120, 290), (760, 290)], fill="#c4a35a", width=2)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _svg_to_png_bytes(diagram_svg: str, width: int = 900) -> bytes | None:
    if not diagram_svg or "<svg" not in diagram_svg:
        return None
    svg = diagram_svg.strip()
    if not svg.startswith("<?xml"):
        svg = '<?xml version="1.0" encoding="UTF-8"?>\n' + svg
    try:
        import cairosvg  # type: ignore

        return cairosvg.svg2png(bytestring=svg.encode("utf-8"), output_width=width)
    except Exception:
        pass
    try:
        from reportlab.graphics import renderPM
        from svglib.svglib import svg2rlg

        with tempfile.TemporaryDirectory() as tmp:
            svg_path = Path(tmp) / "d.svg"
            png_path = Path(tmp) / "d.png"
            svg_path.write_text(svg, encoding="utf-8")
            drawing = svg2rlg(str(svg_path))
            if drawing is None:
                return None
            renderPM.drawToFile(drawing, str(png_path), fmt="PNG")
            return png_path.read_bytes()
    except Exception:
        return None


def diagram_png_for_solution(
    *,
    diagram_svg: str | None,
    diagram_type: str | None = None,
    variables: dict[str, Any] | None = None,
    answers: list[dict[str, Any]] | None = None,
    subject: str | None = None,
) -> bytes:
    """Return PNG bytes for the PDF/DOCX diagram — never empty if we can sketch."""
    png = _svg_to_png_bytes(diagram_svg or "")
    if png:
        return png

    # Try re-render SVG from type
    try:
        from app.diagrams.svg_templates import render_diagram

        params = dict(variables or {})
        for a in answers or []:
            lab = (a.get("label") or "").lower()
            fid = a.get("formula_id") or ""
            if "moment" in lab or "bm" in fid:
                params["M_max"] = a.get("value")
            if "shear" in lab or "sf" in fid:
                params["V_max"] = a.get("value")
            if fid == "rcc.ast_required":
                params["Ast"] = a.get("value")
        dtype = diagram_type or "none"
        if dtype in ("", "none"):
            dtype = _infer_diagram_type(subject, answers, variables)
        svg = render_diagram(dtype, params)
        png = _svg_to_png_bytes(svg or "")
        if png:
            return png
    except Exception:
        pass

    vars_ = variables or {}
    ans_lines = [
        f"{a.get('label')}: {a.get('value')} {a.get('unit')}" for a in (answers or [])[:4]
    ]
    given = [f"{k} = {v}" for k, v in list(vars_.items())[:5]]
    title = {
        "beam_sfd_bmd": "Beam sketch (SFD / BMD)",
        "section": "Cross-section sketch",
        "rcc_section": "RCC section sketch",
        "axial": "Axial member sketch",
    }.get(diagram_type or "", "Problem sketch")
    return _png_from_pillow_fallback(title=title, lines=given + ans_lines)


def _infer_diagram_type(
    subject: str | None,
    answers: list[dict[str, Any]] | None,
    variables: dict[str, Any] | None,
) -> str:
    ids = " ".join((a.get("formula_id") or "") for a in (answers or []))
    keys = " ".join((variables or {}).keys()).lower()
    if any(x in ids for x in ("beam", "moment", "shear", "udl")) or "w" in (variables or {}):
        return "beam_sfd_bmd"
    if "rcc" in ids or (subject or "") == "concrete":
        return "rcc_section"
    if "I" in (variables or {}) or "b" in keys:
        return "section"
    if "P" in (variables or {}) and "A" in (variables or {}):
        return "axial"
    return "beam_sfd_bmd"


def _png_flowable(png: bytes, max_width: float = 150 * mm) -> RLImage:
    bio = io.BytesIO(png)
    img = RLImage(bio)
    img.drawWidth = max_width
    img.drawHeight = max_width * (0.48)  # keep diagram compact so answers fit on page 1
    return img


def _styles():
    base = getSampleStyleSheet()
    return {
        "brand": ParagraphStyle(
            "CMBrand",
            parent=base["Heading1"],
            fontSize=18,
            textColor=INK,
            spaceAfter=2,
            leading=22,
        ),
        "sub": ParagraphStyle(
            "CMSub",
            parent=base["Normal"],
            fontSize=9,
            textColor=MUTED,
            spaceAfter=8,
        ),
        "section": ParagraphStyle(
            "CMSection",
            parent=base["Heading2"],
            fontSize=11,
            textColor=ACCENT,
            spaceBefore=10,
            spaceAfter=6,
            leading=14,
        ),
        "body": ParagraphStyle(
            "CMBody",
            parent=base["BodyText"],
            fontSize=10,
            leading=15,
            textColor=INK,
            spaceAfter=3,
        ),
        "qbody": ParagraphStyle(
            "CMQBody",
            parent=base["BodyText"],
            fontSize=11,
            leading=16,
            textColor=INK,
            spaceAfter=4,
        ),
        "step": ParagraphStyle(
            "CMStep",
            parent=base["BodyText"],
            fontSize=10,
            leading=14,
            textColor=INK,
            leftIndent=4,
        ),
        "answer": ParagraphStyle(
            "CMAnswer",
            parent=base["BodyText"],
            fontSize=12,
            leading=16,
            textColor=A_BORDER,
            alignment=TA_LEFT,
        ),
        "muted": ParagraphStyle(
            "CMMuted",
            parent=base["Normal"],
            fontSize=8,
            textColor=MUTED,
        ),
    }


def _question_block(question: str, styles: dict) -> list[Any]:
    rows = [[Paragraph("<b>QUESTION</b>", styles["body"])]]
    for para in (question or "—").split("\n"):
        if para.strip():
            rows.append([Paragraph(_clean(para), styles["qbody"])])
    t = Table(rows, colWidths=[170 * mm])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), Q_BG),
                ("BOX", (0, 0), (-1, -1), 1.2, ACCENT),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, 0), 8),
                ("BOTTOMPADDING", (0, -1), (-1, -1), 10),
                ("TOPPADDING", (0, 1), (-1, -1), 2),
            ]
        )
    )
    return [t, Spacer(1, 10)]


def _answer_block(answers: list[dict[str, Any]], styles: dict) -> list[Any]:
    if not answers:
        return []
    rows = [[Paragraph("<b>FINAL ANSWER</b>", styles["body"])]]
    for a in answers:
        label = _clean(str(a.get("label") or "Result"))
        val = a.get("value")
        unit = _clean(str(a.get("unit") or ""))
        rows.append(
            [
                Paragraph(
                    f"<b>{label}</b><br/><font size='14'><b>{val}</b> {unit}</font>",
                    styles["answer"],
                )
            ]
        )
    t = Table(rows, colWidths=[170 * mm])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), A_BG),
                ("BOX", (0, 0), (-1, -1), 1.5, A_BORDER),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor("#a7d4bc")),
            ]
        )
    )
    return [Spacer(1, 8), t]


def _solution_steps(explanation: str, steps: list[dict[str, Any]], styles: dict) -> list[Any]:
    out: list[Any] = [_Banner("SOLUTION — step by step", ACCENT), Spacer(1, 8)]

    # Prefer structured registry steps when present
    if steps:
        for i, step in enumerate(steps, start=1):
            name = _clean(str(step.get("name") or f"Step {i}"))
            expr = _clean(str(step.get("expression") or ""))
            inputs = step.get("inputs") or {}
            subst = ", ".join(f"{k}={v}" for k, v in inputs.items())
            val = step.get("value")
            unit = _clean(str(step.get("unit") or ""))
            note = _clean(str(step.get("notes") or ""))
            block = [
                Paragraph(f"<b>Step {i}. {name}</b>", styles["step"]),
                Paragraph(f"Formula: <font face='Courier'>{expr}</font>", styles["step"]),
            ]
            if subst:
                block.append(Paragraph(f"Substitution: {subst}", styles["step"]))
            block.append(Paragraph(f"Result: <b>{val} {unit}</b>", styles["step"]))
            if note:
                block.append(Paragraph(f"Note: {note}", styles["muted"]))
            cell = Table([[b] for b in block], colWidths=[168 * mm])
            cell.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), STEP_BG),
                        ("BOX", (0, 0), (-1, -1), 0.6, LINE),
                        ("LEFTPADDING", (0, 0), (-1, -1), 8),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            out.append(KeepTogether([cell, Spacer(1, 6)]))
        return out

    # Fallback: clean free-text explanation
    for raw in (explanation or "No written steps available.").split("\n"):
        line = _strip_md(raw)
        if not line:
            out.append(Spacer(1, 4))
            continue
        if line.lower().startswith("step") or raw.strip().startswith("##"):
            out.append(Paragraph(f"<b>{_clean(line)}</b>", styles["step"]))
        elif line.startswith("- ") or line.startswith("* "):
            out.append(Paragraph("• " + _clean(line[2:]), styles["step"]))
        else:
            out.append(Paragraph(_clean(line), styles["body"]))
    return out


def build_solution_pdf(
    *,
    title: str,
    question: str,
    explanation: str,
    answers: list[dict[str, Any]],
    diagram_svg: str | None = None,
    subject: str | None = None,
    steps: list[dict[str, Any]] | None = None,
    diagram_type: str | None = None,
    variables: dict[str, Any] | None = None,
    pack_problems: list[dict[str, Any]] | None = None,
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"CivilMaster — {title}",
        author="CivilMaster",
    )
    styles = _styles()
    story: list[Any] = []

    story.append(Paragraph("CivilMaster", styles["brand"]))
    story.append(
        Paragraph(
            "B.Tech Civil | formula-verified worked solution",
            styles["sub"],
        )
    )
    story.append(HRFlowable(width="100%", thickness=1.5, color=ACCENT, spaceAfter=8))
    if title:
        story.append(Paragraph(_clean(title), styles["section"]))
    if subject:
        story.append(
            Paragraph(
                _clean(f"Subject: {str(subject).replace('_', ' ').title()}"),
                styles["muted"],
            )
        )
        story.append(Spacer(1, 6))

    # Multi-question pack: one clear block per question (new page each)
    if pack_problems:
        for idx, prob in enumerate(pack_problems, start=1):
            if idx > 1:
                story.append(PageBreak())
                story.append(Paragraph("CivilMaster", styles["brand"]))
                story.append(
                    Paragraph(
                        "B.Tech Civil | formula-verified worked solution",
                        styles["sub"],
                    )
                )
                story.append(HRFlowable(width="100%", thickness=1.5, color=ACCENT, spaceAfter=8))

            story.append(
                Paragraph(
                    _clean(prob.get("title") or f"Question {idx}"),
                    styles["section"],
                )
            )
            qtext = (prob.get("question") or "")[:1200]
            story.extend(_question_block(qtext, styles))

            kind = (prob.get("kind") or "").lower()
            has_steps = bool(prob.get("steps"))
            has_answers = bool(prob.get("final_answers"))

            if has_steps and has_answers and (
                prob.get("diagram_svg") or prob.get("diagram_type") not in (None, "", "none")
            ):
                story.append(_Banner("DIAGRAM", colors.HexColor("#334155")))
                story.append(Spacer(1, 6))
                png = diagram_png_for_solution(
                    diagram_svg=prob.get("diagram_svg"),
                    diagram_type=prob.get("diagram_type"),
                    variables=prob.get("variables_extracted") or {},
                    answers=prob.get("final_answers") or [],
                    subject=subject,
                )
                story.append(_png_flowable(png))
                story.append(Spacer(1, 8))

            if kind == "detailing":
                story.append(_Banner("GUIDANCE (CAD / detailing)", ACCENT))
                story.append(Spacer(1, 6))
                for line in (prob.get("explanation") or "").split("\n"):
                    line = _strip_md(line)
                    if line:
                        story.append(Paragraph(_clean(line), styles["body"]))
                story.append(Spacer(1, 8))
                story.append(
                    Paragraph(
                        "<i>No fake numerical diagram — draw this in AutoCAD from the guidance above.</i>",
                        styles["muted"],
                    )
                )
            elif has_steps or has_answers:
                story.extend(
                    _solution_steps(
                        prob.get("explanation") or "",
                        prob.get("steps") or [],
                        styles,
                    )
                )
                ans_flow = _answer_block(prob.get("final_answers") or [], styles)
                if ans_flow:
                    story.append(KeepTogether(ans_flow))
            else:
                story.append(_Banner("PARTIAL / NEEDS MANUAL CHECK", colors.HexColor("#9a3412")))
                story.append(Spacer(1, 6))
                for line in (prob.get("explanation") or "").split("\n"):
                    line = _strip_md(line)
                    if line:
                        story.append(Paragraph(_clean(line), styles["body"]))
                story.append(Spacer(1, 6))
                story.append(
                    Paragraph(
                        "Paste this one question alone if you need a single formula-engine check.",
                        styles["muted"],
                    )
                )
        story.append(Spacer(1, 10))
        story.append(
            Paragraph(
                "CivilMaster | Detailing → AutoCAD drawings. Numerical checks → formula engine + IS codes.",
                styles["muted"],
            )
        )
        doc.build(story)
        return buffer.getvalue()

    # Single-problem layout
    q_show = (question or "")[:2000]
    story.extend(_question_block(q_show, styles))

    if steps and (diagram_svg or (diagram_type and diagram_type != "none")):
        story.append(_Banner("DIAGRAM", colors.HexColor("#334155")))
        story.append(Spacer(1, 6))
        png = diagram_png_for_solution(
            diagram_svg=diagram_svg,
            diagram_type=diagram_type,
            variables=variables,
            answers=answers,
            subject=subject,
        )
        story.append(_png_flowable(png))
        story.append(Spacer(1, 10))

    story.extend(_solution_steps(explanation, steps or [], styles))

    answers_flow = _answer_block(answers, styles)
    if answers_flow:
        story.append(KeepTogether(answers_flow))

    story.append(Spacer(1, 10))
    story.append(
        Paragraph(
            "CivilMaster | cross-check with IS codes and faculty guidance before submission.",
            styles["muted"],
        )
    )
    doc.build(story)
    return buffer.getvalue()


def build_solution_docx(
    *,
    title: str,
    question: str,
    explanation: str,
    answers: list[dict[str, Any]],
    subject: str | None = None,
    diagram_svg: str | None = None,
    steps: list[dict[str, Any]] | None = None,
    diagram_type: str | None = None,
    variables: dict[str, Any] | None = None,
) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt, RGBColor

    def shade(paragraph, fill: str):
        p = paragraph._p
        pPr = p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        shd.set(qn("w:val"), "clear")
        pPr.append(shd)

    doc = Document()
    h = doc.add_heading("CivilMaster", level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    sub = doc.add_paragraph("B.Tech Civil · formula-verified worked solution")
    for run in sub.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    if title:
        doc.add_heading(title, level=1)
    if subject:
        doc.add_paragraph(f"Subject: {subject.replace('_', ' ').title()}")

    qh = doc.add_paragraph()
    run = qh.add_run("QUESTION")
    run.bold = True
    run.font.color.rgb = RGBColor(0x1D, 0x4E, 0x89)
    shade(qh, "EEF4FB")
    qp = doc.add_paragraph(question or "")
    shade(qp, "EEF4FB")

    dh = doc.add_paragraph()
    run = dh.add_run("DIAGRAM")
    run.bold = True
    png = diagram_png_for_solution(
        diagram_svg=diagram_svg,
        diagram_type=diagram_type,
        variables=variables,
        answers=answers,
        subject=subject,
    )
    doc.add_picture(io.BytesIO(png), width=Inches(5.9))

    sh = doc.add_paragraph()
    run = sh.add_run("SOLUTION — step by step")
    run.bold = True
    run.font.color.rgb = RGBColor(0x1D, 0x4E, 0x89)

    if steps:
        for i, step in enumerate(steps, start=1):
            doc.add_paragraph(
                f"Step {i}. {step.get('name')}",
                style="List Number" if False else None,
            )
            p = doc.add_paragraph()
            p.add_run("Formula: ").bold = True
            p.add_run(str(step.get("expression") or ""))
            inputs = step.get("inputs") or {}
            if inputs:
                doc.add_paragraph(
                    "Substitution: " + ", ".join(f"{k}={v}" for k, v in inputs.items())
                )
            r = doc.add_paragraph()
            r.add_run("Result: ").bold = True
            r.add_run(f"{step.get('value')} {step.get('unit') or ''}")
    else:
        for line in (explanation or "No written steps available.").split("\n"):
            line = _strip_md(line)
            if line:
                doc.add_paragraph(line)

    if answers:
        ah = doc.add_paragraph()
        run = ah.add_run("FINAL ANSWER")
        run.bold = True
        run.font.color.rgb = RGBColor(0x2F, 0x6B, 0x4F)
        shade(ah, "E8F5EF")
        for a in answers:
            ap = doc.add_paragraph()
            ap.add_run(f"{a.get('label')}: ").bold = True
            ap.add_run(f"{a.get('value')} {a.get('unit') or ''}")
            shade(ap, "E8F5EF")

    note = doc.add_paragraph(
        "CivilMaster · cross-check with IS codes and faculty guidance before submission."
    )
    for run in note.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build_beam_dxf(variables: dict[str, Any] | None = None) -> bytes:
    import ezdxf

    variables = variables or {}
    length = float(variables.get("L") or variables.get("span") or 4000.0)
    scale = 200.0 / max(length, 1.0)
    L = length * scale

    doc = ezdxf.new("R2010")
    msp = doc.modelspace()
    msp.add_line((0, 0), (L, 0))
    msp.add_line((0, 0), (-5, -12))
    msp.add_line((0, 0), (5, -12))
    msp.add_line((L, 0), (L - 5, -12))
    msp.add_line((L, 0), (L + 5, -12))
    msp.add_line((0, 25), (L, 25))
    step = max(10, int(L // 20) or 10)
    for i in range(0, int(L) + 1, step):
        msp.add_line((i, 25), (i, 12))
    t1 = msp.add_text("Simply supported beam (CivilMaster)", dxfattribs={"height": 8})
    t1.dxf.insert = (0, 40)
    t2 = msp.add_text(f"L = {length}", dxfattribs={"height": 5})
    t2.dxf.insert = (0, -28)

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "beam.dxf"
        doc.saveas(path)
        return path.read_bytes()


def build_section_dxf(variables: dict[str, Any] | None = None) -> bytes:
    import ezdxf

    variables = variables or {}
    b = float(variables.get("b") or 200.0)
    d = float(variables.get("d") or 300.0)
    scale = 120.0 / max(b, d, 1.0)
    W, H = b * scale, d * scale

    doc = ezdxf.new("R2010")
    msp = doc.modelspace()
    msp.add_lwpolyline([(0, 0), (W, 0), (W, H), (0, H), (0, 0)], close=True)
    cx, cy = W / 2, H / 2
    msp.add_line((cx - 6, cy), (cx + 6, cy))
    msp.add_line((cx, cy - 6), (cx, cy + 6))
    t = msp.add_text(f"Section b={b} d={d} (CivilMaster)", dxfattribs={"height": 6})
    t.dxf.insert = (0, H + 12)

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "section.dxf"
        doc.saveas(path)
        return path.read_bytes()


def build_rcc_section_dxf(variables: dict[str, Any] | None = None) -> bytes:
    import ezdxf

    variables = variables or {}
    b = float(variables.get("b") or 230.0)
    d = float(variables.get("d") or 450.0)
    scale = 140.0 / max(b, d, 1.0)
    W, H = b * scale, d * scale

    doc = ezdxf.new("R2010")
    msp = doc.modelspace()
    msp.add_lwpolyline([(0, 0), (W, 0), (W, H), (0, H), (0, 0)], close=True)
    y_bar = 12
    for x in (W * 0.25, W * 0.5, W * 0.75):
        msp.add_circle((x, y_bar), 4)
    t = msp.add_text("RCC section (CivilMaster)", dxfattribs={"height": 6})
    t.dxf.insert = (0, H + 12)

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "rcc.dxf"
        doc.saveas(path)
        return path.read_bytes()


def build_solution_dxf(
    *,
    diagram_type: str | None = None,
    variables: dict[str, Any] | None = None,
) -> bytes:
    dtype = (diagram_type or "").lower()
    if dtype in ("section", "rect_section"):
        return build_section_dxf(variables)
    if dtype in ("rcc_section", "rcc"):
        return build_rcc_section_dxf(variables)
    return build_beam_dxf(variables)
